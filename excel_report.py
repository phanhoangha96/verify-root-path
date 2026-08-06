from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from mapping import EXPECTED_ROOT_PATH_FOLDERS, expected_folders_for


@dataclass
class AuditResult:
    attachment_id: str
    file_name: Optional[str]
    file_path: Optional[str]
    root_path: Optional[str]
    expected_folders: str
    found_server_path: Optional[str]
    matched_server_root: str
    status: str
    note: str


DETAIL_HEADERS = [
    "ATTACHMENT_ID",
    "FILE_NAME",
    "FILE_PATH",
    "ROOT_PATH",
    "EXPECTED_FOLDERS",
    "FOUND_SERVER_PATH",
    "MATCHED_SERVER_ROOT",
    "STATUS",
    "NOTE",
]


def write_excel(results: List[AuditResult], output_path: Path) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook(write_only=True)

    summary_ws = wb.create_sheet("Summary_by_ROOT_PATH")
    summary_ws.append(
        [
            "ROOT_PATH",
            "EXPECTED_FOLDERS",
            "TOTAL_SAMPLED",
            "FOUND",
            "NOT_FOUND",
            "MATCHED",
            "MISMATCH",
            "DISCOVERED_SERVER_ROOTS",
            "DISCOVERED_ROOT_COUNTS",
        ]
    )

    detail_ws = wb.create_sheet("Detail")
    detail_ws.append(DETAIL_HEADERS)

    no_root_ws = wb.create_sheet("No_ROOT_PATH")
    no_root_ws.append(DETAIL_HEADERS)

    mismatch_ws = wb.create_sheet("Mismatch")
    mismatch_ws.append(DETAIL_HEADERS)

    mapping_ws = wb.create_sheet("Expected_Mapping")
    mapping_ws.append(["ROOT_PATH", "EXPECTED_FOLDER"])
    for key, folders in EXPECTED_ROOT_PATH_FOLDERS.items():
        for folder in folders:
            mapping_ws.append([key, folder])

    # Aggregate
    group_total: Counter = Counter()
    group_found: Counter = Counter()
    group_not_found: Counter = Counter()
    group_matched: Counter = Counter()
    group_mismatch: Counter = Counter()
    discovered: Dict[str, Counter] = defaultdict(Counter)

    for row in results:
        key = row.root_path or "(NULL)"
        group_total[key] += 1
        values = [
            row.attachment_id,
            row.file_name or "",
            row.file_path or "",
            row.root_path or "",
            row.expected_folders,
            row.found_server_path or "",
            row.matched_server_root,
            row.status,
            row.note,
        ]
        detail_ws.append(values)

        if not row.root_path:
            no_root_ws.append(values)
        if row.status in {"MISMATCH", "UNKNOWN_ROOT_PATH_FOUND"}:
            mismatch_ws.append(values)

        if row.found_server_path:
            group_found[key] += 1
            if row.matched_server_root:
                discovered[key][row.matched_server_root] += 1
        else:
            group_not_found[key] += 1

        if row.status == "MATCHED":
            group_matched[key] += 1
        elif row.status == "MISMATCH":
            group_mismatch[key] += 1

    # Stable key order: known mapping keys first, then others, NULL last
    known = list(EXPECTED_ROOT_PATH_FOLDERS.keys())
    other_keys = sorted(k for k in group_total if k not in known and k != "(NULL)")
    ordered_keys = known + other_keys
    if "(NULL)" in group_total:
        ordered_keys.append("(NULL)")

    for key in ordered_keys:
        if key not in group_total:
            continue
        expected = ", ".join(expected_folders_for(None if key == "(NULL)" else key))
        disc = discovered.get(key, Counter())
        disc_roots = ", ".join(sorted(disc.keys()))
        disc_counts = ", ".join(f"{root}={count}" for root, count in sorted(disc.items()))
        summary_ws.append(
            [
                key,
                expected,
                group_total[key],
                group_found[key],
                group_not_found[key],
                group_matched[key],
                group_mismatch[key],
                disc_roots,
                disc_counts,
            ]
        )

    overview = wb.create_sheet("Overview", 0)
    overview.append(["Metric", "Value"])
    overview.append(["Generated_at", datetime.now().isoformat(timespec="seconds")])
    overview.append(["Total_rows", len(results)])
    overview.append(["No_ROOT_PATH", sum(1 for r in results if not r.root_path)])
    overview.append(["Found", sum(1 for r in results if r.found_server_path)])
    overview.append(["Not_found", sum(1 for r in results if not r.found_server_path)])
    overview.append(["Matched", sum(1 for r in results if r.status == "MATCHED")])
    overview.append(["Mismatch", sum(1 for r in results if r.status == "MISMATCH")])
    status_counts = Counter(r.status for r in results)
    overview.append([])
    overview.append(["Status", "Count"])
    for status, count in sorted(status_counts.items()):
        overview.append([status, count])

    wb.save(output_path)
    return output_path
