"""Write missing-files report (CSV always; XLSX when row count fits Excel limit)."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, Iterable, Iterator, List, Optional, Tuple

from openpyxl import Workbook

from db import _split_csv_line, csv_escape

EXCEL_MAX_DATA_ROWS = 1_048_575

# Excel/XML disallow most C0 control chars (openpyxl raises IllegalCharacterError).
# Keep tab/LF/CR; strip the rest.
_ILLEGAL_EXCEL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


@dataclass
class ReportSummary:
    """Aggregated stats written to the Excel Summary sheet (and used for CSV header notes)."""

    total_input_rows: int = 0
    unique_files: int = 0
    missing_files: int = 0
    found_files: int = 0
    duplicate_input_skipped: int = 0
    elapsed_seconds: float = 0.0
    check_rate_per_sec: float = 0.0
    started_at: str = ""
    finished_at: str = ""
    job_dir: str = ""
    input_csv: str = ""
    missing_csv: str = ""
    extra: Dict[str, str] = field(default_factory=dict)


def sanitize_excel_text(value: object) -> str:
    """Remove characters that Excel worksheets cannot store."""
    if value is None:
        return ""
    text = str(value)
    if not text:
        return ""
    return _ILLEGAL_EXCEL_CHARS.sub("", text)


def iter_unique_path_filename(csv_file: Path) -> Iterator[Tuple[str, str]]:
    """Yield unique (path, fileName) pairs from a path,fileName CSV (first occurrence wins)."""
    seen: set[Tuple[str, str]] = set()
    with csv_file.open("r", encoding="utf-8") as reader:
        header = reader.readline()
        if not header:
            return
        for line in reader:
            if not line.strip():
                continue
            parts = _split_csv_line(line.rstrip("\n\r"))
            path = parts[0] if parts else ""
            name = parts[1] if len(parts) > 1 else ""
            key = (path, name)
            if key in seen:
                continue
            seen.add(key)
            yield path, name


def count_unique_path_filename(csv_file: Path) -> Tuple[int, int]:
    """Return (unique_count, raw_data_rows) for a path,fileName CSV."""
    unique = 0
    raw = 0
    seen: set[Tuple[str, str]] = set()
    with csv_file.open("r", encoding="utf-8") as reader:
        header = reader.readline()
        if not header:
            return 0, 0
        for line in reader:
            if not line.strip():
                continue
            raw += 1
            parts = _split_csv_line(line.rstrip("\n\r"))
            key = (parts[0] if parts else "", parts[1] if len(parts) > 1 else "")
            if key in seen:
                continue
            seen.add(key)
            unique += 1
    return unique, raw


def format_duration(seconds: float) -> str:
    total = int(max(0.0, seconds))
    hours, rem = divmod(total, 3600)
    minutes, secs = divmod(rem, 60)
    if hours:
        return f"{hours}h{minutes:02d}m{secs:02d}s"
    if minutes:
        return f"{minutes}m{secs:02d}s"
    millis = int(round((max(0.0, seconds) - total) * 1000))
    if millis and total < 10:
        return f"{total}.{millis:03d}s"
    return f"{secs}s"


def write_missing_excel(
    missing_csv: Path,
    output_xlsx: Path,
    *,
    summary: Optional[ReportSummary] = None,
) -> Path:
    """
    Write XLSX with:
      - Summary: totals, duration, paths
      - Missing files: unique (path, fileName) only
    """
    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    unique_missing, raw_missing = count_unique_path_filename(missing_csv)
    if summary is None:
        summary = ReportSummary(missing_files=unique_missing)
    else:
        summary.missing_files = unique_missing

    if unique_missing > EXCEL_MAX_DATA_ROWS:
        raise SystemExit(
            f"Too many unique missing rows for Excel ({unique_missing}). "
            f"Excel max is {EXCEL_MAX_DATA_ROWS}. Use CSV output instead."
        )

    wb = Workbook(write_only=True)
    _write_summary_sheet(wb, summary, raw_missing_rows=raw_missing)
    sheet = wb.create_sheet("Missing files")
    sheet.append(["STT", "Path", "File name"])

    row_index = 1
    for path, name in iter_unique_path_filename(missing_csv):
        sheet.append(
            [
                row_index,
                sanitize_excel_text(path),
                sanitize_excel_text(name),
            ]
        )
        row_index += 1

    wb.save(output_xlsx)
    return output_xlsx


def write_missing_csv_unique(
    missing_csv: Path,
    output_csv: Path,
    *,
    summary: Optional[ReportSummary] = None,
) -> Path:
    """Copy missing rows to output CSV, deduped by (path, fileName)."""
    output_csv.parent.mkdir(parents=True, exist_ok=True)
    with output_csv.open("w", encoding="utf-8", newline="") as writer:
        if summary is not None:
            for line in _summary_comment_lines(summary):
                writer.write(f"# {line}\n")
        writer.write("path,fileName\n")
        for path, name in iter_unique_path_filename(missing_csv):
            writer.write(f"{csv_escape(path)},{csv_escape(name)}\n")
    return output_csv


def _write_summary_sheet(
    wb: Workbook,
    summary: ReportSummary,
    *,
    raw_missing_rows: int,
) -> None:
    sheet = wb.create_sheet("Summary")
    sheet.append(["Metric", "Value"])

    unique_files = summary.unique_files
    missing = summary.missing_files
    found = summary.found_files
    if unique_files > 0 and found <= 0 and missing >= 0:
        found = max(0, unique_files - missing)

    rows: List[Tuple[str, object]] = [
        ("Total input rows (CSV)", summary.total_input_rows or ""),
        ("Unique files (path + fileName)", unique_files or ""),
        ("Duplicate input rows skipped", summary.duplicate_input_skipped or 0),
        ("Files found", found if unique_files else ""),
        ("Files missing (unique)", missing),
        (
            "Missing ratio",
            f"{(missing / unique_files * 100):.2f}%" if unique_files else "",
        ),
        ("Raw missing.csv rows (before dedupe)", raw_missing_rows),
        ("Elapsed", format_duration(summary.elapsed_seconds) if summary.elapsed_seconds else ""),
        (
            "Check rate",
            f"{summary.check_rate_per_sec:.1f}/s" if summary.check_rate_per_sec else "",
        ),
        ("Started at", summary.started_at),
        ("Finished at", summary.finished_at),
        ("Job dir", summary.job_dir),
        ("Input CSV", summary.input_csv),
        ("Missing CSV", summary.missing_csv),
    ]
    for key, value in summary.extra.items():
        rows.append((key, value))

    for metric, value in rows:
        sheet.append([sanitize_excel_text(metric), sanitize_excel_text(value)])


def _summary_comment_lines(summary: ReportSummary) -> Iterable[str]:
    unique_files = summary.unique_files
    missing = summary.missing_files
    found = summary.found_files
    if unique_files > 0 and found <= 0:
        found = max(0, unique_files - missing)
    yield f"unique_files={unique_files}"
    yield f"missing_files={missing}"
    yield f"found_files={found}"
    yield f"duplicate_input_skipped={summary.duplicate_input_skipped}"
    yield f"total_input_rows={summary.total_input_rows}"
    if summary.elapsed_seconds:
        yield f"elapsed={format_duration(summary.elapsed_seconds)}"
    if summary.check_rate_per_sec:
        yield f"rate={summary.check_rate_per_sec:.1f}/s"


def choose_output_format(missing_count: int, forced: str = "") -> str:
    """Return 'xlsx' or 'csv'. Auto: xlsx if missing fits Excel sheet, else csv."""
    forced = (forced or "").strip().lower()
    if forced in {"csv", "xlsx", "excel"}:
        return "csv" if forced == "csv" else "xlsx"
    return "csv" if missing_count > EXCEL_MAX_DATA_ROWS else "xlsx"


def read_missing_preview(missing_csv: Path, limit: int = 5) -> List[Tuple[str, str]]:
    rows: List[Tuple[str, str]] = []
    for path, name in iter_unique_path_filename(missing_csv):
        rows.append((path, name))
        if len(rows) >= limit:
            break
    return rows
