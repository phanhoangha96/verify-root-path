from __future__ import annotations

import json
from collections import Counter
from dataclasses import asdict, dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Dict, List

from openpyxl import Workbook

MAX_DETAIL_ROWS = 20000


@dataclass
class IssueRow:
    source: str
    object_type: int
    doc_id: str
    reason: str
    detail: str = ""


@dataclass
class BackfillReport:
    started_at: str
    finished_at: str = ""
    duration_seconds: float = 0.0
    dry_run: bool = False
    source: str = ""
    object_type: str = ""
    solr_host: str = ""
    solr_core: str = ""
    scanned: int = 0
    indexed: int = 0
    skipped_missing_dept: int = 0
    skipped_empty_search_text: int = 0
    errors: int = 0
    solr_batches: int = 0
    by_key: Dict[str, Dict[str, int]] = field(default_factory=dict)
    skipped_rows: List[IssueRow] = field(default_factory=list)
    error_rows: List[IssueRow] = field(default_factory=list)

    def bump(self, source: str, object_type: int, field: str, count: int = 1) -> None:
        key = f"{source}:{object_type}"
        bucket = self.by_key.setdefault(
            key,
            {
                "source": source,
                "object_type": object_type,
                "scanned": 0,
                "indexed": 0,
                "skipped_missing_dept": 0,
                "skipped_empty_search_text": 0,
                "errors": 0,
            },
        )
        bucket[field] = int(bucket.get(field, 0)) + count
        setattr(self, field, int(getattr(self, field)) + count)

    def add_skip(self, row: IssueRow) -> None:
        field = (
            "skipped_missing_dept"
            if row.reason == "MISSING_DEPT_ID"
            else "skipped_empty_search_text"
        )
        self.bump(row.source, row.object_type, field)
        if len(self.skipped_rows) < MAX_DETAIL_ROWS:
            self.skipped_rows.append(row)

    def add_error(self, row: IssueRow) -> None:
        self.bump(row.source, row.object_type, "errors")
        if len(self.error_rows) < MAX_DETAIL_ROWS:
            self.error_rows.append(row)


def write_reports(report: BackfillReport, output_dir: Path, stamp: str, output_file: str = "") -> Dict[str, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    if output_file:
        xlsx_path = Path(output_file)
        if not xlsx_path.is_absolute():
            xlsx_path = output_dir / output_file
        json_path = xlsx_path.with_suffix(".json")
    else:
        base = f"solr_doc_meta_backfill_{stamp}"
        xlsx_path = output_dir / f"{base}.xlsx"
        json_path = output_dir / f"{base}.json"

    payload = {
        "started_at": report.started_at,
        "finished_at": report.finished_at,
        "duration_seconds": report.duration_seconds,
        "dry_run": report.dry_run,
        "source": report.source,
        "object_type": report.object_type,
        "solr_host": report.solr_host,
        "solr_core": report.solr_core,
        "scanned": report.scanned,
        "indexed": report.indexed,
        "skipped_missing_dept": report.skipped_missing_dept,
        "skipped_empty_search_text": report.skipped_empty_search_text,
        "errors": report.errors,
        "solr_batches": report.solr_batches,
        "by_key": report.by_key,
        "skipped_truncated": len(report.skipped_rows) >= MAX_DETAIL_ROWS,
        "errors_truncated": len(report.error_rows) >= MAX_DETAIL_ROWS,
        "skipped": [asdict(row) for row in report.skipped_rows],
        "error_rows": [asdict(row) for row in report.error_rows],
    }
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    wb = Workbook(write_only=True)
    overview = wb.create_sheet("Overview")
    overview.append(["Metric", "Value"])
    for key, value in [
        ("Generated_at", datetime.now().isoformat(timespec="seconds")),
        ("Started_at", report.started_at),
        ("Finished_at", report.finished_at),
        ("Duration_seconds", report.duration_seconds),
        ("Dry_run", report.dry_run),
        ("Source", report.source),
        ("Object_type", report.object_type),
        ("Solr_host", report.solr_host),
        ("Solr_core", report.solr_core),
        ("Scanned", report.scanned),
        ("Indexed", report.indexed),
        ("Skipped_missing_dept", report.skipped_missing_dept),
        ("Skipped_empty_search_text", report.skipped_empty_search_text),
        ("Errors", report.errors),
        ("Solr_batches", report.solr_batches),
    ]:
        overview.append([key, value])

    summary = wb.create_sheet("By_source")
    summary.append(
        [
            "SOURCE",
            "OBJECT_TYPE",
            "SCANNED",
            "INDEXED",
            "SKIPPED_MISSING_DEPT",
            "SKIPPED_EMPTY_SEARCH_TEXT",
            "ERRORS",
        ]
    )
    for bucket in report.by_key.values():
        summary.append(
            [
                bucket.get("source"),
                bucket.get("object_type"),
                bucket.get("scanned", 0),
                bucket.get("indexed", 0),
                bucket.get("skipped_missing_dept", 0),
                bucket.get("skipped_empty_search_text", 0),
                bucket.get("errors", 0),
            ]
        )

    skipped = wb.create_sheet("Skipped")
    skipped.append(["SOURCE", "OBJECT_TYPE", "DOC_ID", "REASON", "DETAIL"])
    for row in report.skipped_rows:
        skipped.append([row.source, row.object_type, row.doc_id, row.reason, row.detail])

    errors = wb.create_sheet("Errors")
    errors.append(["SOURCE", "OBJECT_TYPE", "DOC_ID", "REASON", "DETAIL"])
    for row in report.error_rows:
        errors.append([row.source, row.object_type, row.doc_id, row.reason, row.detail])

    reasons = wb.create_sheet("Skip_reasons")
    reasons.append(["REASON", "COUNT"])
    for reason, count in sorted(Counter(r.reason for r in report.skipped_rows).items()):
        reasons.append([reason, count])

    wb.save(xlsx_path)
    return {"xlsx": xlsx_path, "json": json_path}
